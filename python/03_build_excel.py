#!/usr/bin/env python3
"""Data dictionary, cleaning log, KPI reconciliation workbooks."""
from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows

ROOT = Path(__file__).resolve().parents[1]
EXCEL = ROOT / "excel"
MARTS = ROOT / "data" / "marts"
REPORTS = ROOT / "reports"
EXCEL.mkdir(parents=True, exist_ok=True)

HEADER = Font(bold=True, color="1B1A19")
HFILL = PatternFill("solid", fgColor="742774")
THIN = Border(
    left=Side(style="thin", color="CBD5E1"),
    right=Side(style="thin", color="CBD5E1"),
    top=Side(style="thin", color="CBD5E1"),
    bottom=Side(style="thin", color="CBD5E1"),
)


def style_header(ws):
    for cell in ws[1]:
        cell.font = HEADER
        cell.fill = HFILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = THIN


def autosize(ws, widths=None):
    for i, col in enumerate(ws.columns, 1):
        letter = col[0].column_letter
        if widths and i <= len(widths):
            ws.column_dimensions[letter].width = widths[i - 1]
        else:
            ws.column_dimensions[letter].width = min(42, max(12, max(len(str(c.value or "")) for c in col) + 2))


def write_df(ws, df):
    for r in dataframe_to_rows(df, index=False, header=True):
        ws.append(r)
    style_header(ws)


def build_dictionary():
    rows = [
        ("dim_store", "store_id", "string", "PK", "Dark store identifier"),
        ("dim_store", "metro", "string", "", "India metro cluster"),
        ("dim_store", "zone", "string", "", "Micro-zone / catchment"),
        ("dim_store", "promise_minutes", "int", "", "Customer promise SLA (minutes)"),
        ("dim_sku", "sku_id", "string", "PK", "SKU identifier"),
        ("dim_sku", "category", "string", "", "Merch category"),
        ("dim_sku", "mrp", "float", "INR", "List price"),
        ("dim_sku", "perishable", "bool", "", "Fresh / dairy / frozen flag"),
        ("dim_rider", "rider_id", "string", "PK", "Rider identifier"),
        ("dim_rider", "store_id", "string", "FK", "Home dark store"),
        ("fact_orders", "order_id", "string", "PK", "Order header"),
        ("fact_orders", "ordered_at", "datetime", "", "Order placed timestamp"),
        ("fact_orders", "o2d_minutes", "float", "min", "Order-to-door time (delivered only)"),
        ("fact_orders", "sla_hit", "int", "0/1", "1 if o2d <= promise_minutes"),
        ("fact_orders", "gmv", "float", "INR", "Merchandise GMV (fulfilled lines)"),
        ("fact_orders", "net_gmv", "float", "INR", "GMV + delivery fee − discount"),
        ("fact_orders", "had_stockout", "int", "0/1", "Any line stocked out on this order"),
        ("fact_orders", "promo_type", "string", "", "Promo mechanic applied"),
        ("fact_order_lines", "order_id", "string", "FK", "Parent order"),
        ("fact_order_lines", "qty_fulfilled", "int", "", "Units picked; 0 on stockout"),
        ("fact_stockouts", "event_id", "string", "PK", "Stockout event"),
        ("fact_stockouts", "reason", "string", "", "Zero on hand / pick fail / damaged / expired"),
        ("mart_daily", "sla_hit_pct", "float", "", "Delivered SLA hit rate for the day"),
        ("mart_store", "stockout_rate", "float", "", "Orders with ≥1 stockout / orders"),
        ("mart_promo", "promo_roi_proxy", "float", "", "Rough ROI vs no-promo baseline AOV"),
    ]
    df = pd.DataFrame(rows, columns=["table", "column", "type", "key_or_unit", "description"])
    wb = Workbook()
    ws = wb.active
    ws.title = "dictionary"
    write_df(ws, df)
    autosize(ws, [16, 18, 10, 12, 48])
    ws2 = wb.create_sheet("grain_notes")
    ws2.append(["table", "grain", "notes"])
    style_header(ws2)
    for r in [
        ("fact_orders", "1 row = 1 order", "Cancelled orders keep header with GMV=0"),
        ("fact_order_lines", "1 row = 1 SKU on an order", "Includes stocked-out lines (qty_fulfilled=0)"),
        ("mart_daily", "1 row = 1 calendar day", "Built from mart fact_orders"),
        ("mart_store", "1 row = 1 dark store", "SLA / cancel / stockout rolled up"),
    ]:
        ws2.append(list(r))
    autosize(ws2)
    wb.save(EXCEL / "01_data_dictionary.xlsx")


def build_cleaning_log():
    wb = Workbook()
    ws = wb.active
    ws.title = "cleaning_log"
    rows = [
        ["step", "layer", "action", "rows_in", "rows_out", "notes"],
        ["1", "landing", "Land ops extract CSVs into data/landing", "—", "—", "Stores, SKUs, riders, orders, lines, stockouts"],
        ["2", "cleaned", "Parse timestamps; add is_delivered / is_cancelled", "47959", "47959", "Partial treated as delivered for GMV"],
        ["3", "cleaned", "Assert order_id unique; gmv >= 0", "47959", "47959", "No drops required on this extract"],
        ["4", "marts", "Build dim_date + fact_orders with gross_margin", "47959", "47959", "gross_margin = net_gmv − cogs"],
        ["5", "marts", "Aggregate mart_daily / store / zone / promo / rider", "—", "—", "Percentiles via numpy on delivered O2D"],
        ["6", "recon", "Python KPI pack ≡ SQL 04_kpi_queries", "—", "—", "See 03_kpi_reconciliation.xlsx"],
    ]
    for r in rows:
        ws.append(r)
    style_header(ws)
    autosize(ws, [8, 10, 48, 10, 10, 48])
    ws2 = wb.create_sheet("open_questions")
    ws2.append(["item", "status", "note"])
    style_header(ws2)
    for r in [
        ("Guest / logged-out identity", "out of scope", "Extract has no customer_id — CRM deferred"),
        ("Promise clock start", "assumed order_placed", "Confirm if promise starts at payment confirm"),
        ("Rider break / offline time", "approx", "Utilization uses fixed shift hours × 2.5 orders/hr"),
    ]:
        ws2.append(list(r))
    autosize(ws2)
    wb.save(EXCEL / "02_cleaning_log.xlsx")


def build_recon():
    kpi = json.loads((REPORTS / "kpi_summary.json").read_text())
    fact = pd.read_csv(MARTS / "fact_orders.csv")
    delivered = fact[fact["is_delivered"] == 1]

    checks = [
        ("orders", kpi["orders"], int(len(fact)), "count fact_orders"),
        ("delivered_orders", kpi["delivered_orders"], int(delivered.shape[0]), "is_delivered=1"),
        ("gmv_inr", kpi["gmv_inr"], round(float(delivered["gmv"].sum()), 2), "sum gmv delivered"),
        ("aov_inr", kpi["aov_inr"], round(float(delivered["gmv"].sum() / len(delivered)), 2), "gmv / delivered"),
        ("sla_hit_pct", kpi["sla_hit_pct"], round(float(delivered["sla_hit"].mean()), 4), "mean sla_hit delivered"),
        ("cancel_pct", kpi["cancel_pct"], round(float(fact["is_cancelled"].mean()), 4), "mean is_cancelled"),
        ("stockout_rate", kpi["stockout_rate"], round(float(fact["had_stockout"].mean()), 4), "mean had_stockout"),
        ("o2d_p50_min", kpi["o2d_p50_min"], round(float(delivered["o2d_minutes"].median()), 2), "median o2d delivered"),
    ]
    wb = Workbook()
    ws = wb.active
    ws.title = "kpi_recon"
    ws.append(["kpi", "json_value", "recomputed", "delta", "logic"])
    style_header(ws)
    for name, jv, rv, logic in checks:
        delta = round(float(jv) - float(rv), 6) if isinstance(jv, float) else int(jv) - int(rv)
        ws.append([name, jv, rv, delta, logic])
    autosize(ws, [18, 16, 16, 12, 36])

    ws2 = wb.create_sheet("layer_counts")
    ws2.append(["layer", "file", "rows"])
    style_header(ws2)
    for layer, folder in [("landing", ROOT / "data" / "landing"), ("cleaned", ROOT / "data" / "cleaned"), ("marts", MARTS)]:
        for f in sorted(folder.glob("*.csv")):
            n = sum(1 for _ in open(f, encoding="utf-8")) - 1
            ws2.append([layer, f.name, n])
    autosize(ws2)
    wb.save(EXCEL / "03_kpi_reconciliation.xlsx")


if __name__ == "__main__":
    build_dictionary()
    build_cleaning_log()
    build_recon()
    print("Excel workbooks written to excel/")
